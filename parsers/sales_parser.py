"""
Massage Envy — Gunn Group FDA/Manager Sales Dashboard: parsing.
Extracted from the original sales_dashboard.py CLI script. The web app
supplies `prev` (the most recently saved sales report, if any) instead of
reading sales_dashboard_prev.json from disk.
"""

import re
from collections import defaultdict
from datetime import datetime
from openpyxl import load_workbook

TARGET_CATS    = ("Front Desk", "Manager")
POOR_THRESHOLD = 0.1501
LTV            = 80 * 18   # $80/month × 18 months

CLINIC_KEYS = {
    'Gleannloch Farms (0131)', 'Market Street (0124)', 'The Woodlands (0059)',
    'West Katy-Firethorne (1223)', 'Copperfield (0106)', 'Cypress (0157)', 'Katy (0040)'
}
CLINIC_ORDER = [
    'Gleannloch Farms (0131)', 'Market Street (0124)', 'The Woodlands (0059)',
    'West Katy-Firethorne (1223)', 'Copperfield (0106)', 'Cypress (0157)', 'Katy (0040)'
]


def parse_report(input_path, prev):
    """
    prev: dict shaped like the old sales_dashboard_prev.json --
          {"month_label": ..., "run_date": ..., "employees": {key: {guests,memb,close}}, "clinic_ltv": {...}}
          or {} if there's no prior run at all.
    """
    wb = load_workbook(input_path, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))

    month_label = "Unknown Period"
    date_range  = ""
    for r in rows[:10]:
        cell = str(r[1]) if len(r) > 1 and r[1] else ""
        if "Date Range:" in cell:
            match = re.search(r"(\d+/\d+/\d+) to (\d+/\d+/\d+)", cell)
            if match:
                start = datetime.strptime(match.group(1), "%m/%d/%Y")
                end   = datetime.strptime(match.group(2), "%m/%d/%Y")
                month_label = start.strftime("%B %Y")
                date_range  = f"{start.strftime('%m/%d/%Y')} – {end.strftime('%m/%d/%Y')}"
            break

    prev_month      = prev.get("month_label", "")
    prev_run_date   = prev.get("run_date", "")
    same_month      = (prev_month == month_label)
    prev_emps       = prev.get("employees", {})
    prev_clinic_ltv = prev.get("clinic_ltv", {})

    current_clinic = None
    employees = []

    for r in rows:
        if len(r) < 17: continue
        if r[3] and str(r[3]) in CLINIC_KEYS: current_clinic = r[3]
        if not r[4] or not r[5] or r[5] == 'Employee Category': continue
        if not any(t in str(r[5]) for t in TARGET_CATS): continue
        if not current_clinic: continue

        guests = int(r[9] or 0)
        memb   = int(r[10] or 0)
        close  = float(r[11]) if r[11] and r[11] != 'n/a' else 0.0
        retail = float(r[13] or 0)
        gc     = float(r[16] or 0)

        key = f"{current_clinic}|{r[4]}"
        prev_data = prev_emps.get(key)

        employees.append({
            "key":         key,
            "name":        r[4],
            "clinic":      current_clinic,
            "category":    r[5],
            "guests":      guests,
            "memb":        memb,
            "close":       close,
            "retail":      retail,
            "gc":          gc,
            "is_new":      prev_data is None and bool(prev_emps),
            "prev_guests": prev_data["guests"] if prev_data else None,
            "prev_memb":   prev_data["memb"]   if prev_data else None,
            "prev_close":  prev_data["close"]  if prev_data else None,
            "same_month":  same_month,
        })

    employees.sort(key=lambda x: -x['close'])
    return employees, month_label, date_range, prev_month, prev_run_date, prev_clinic_ltv


def peek_report_month(input_path):
    wb = load_workbook(input_path, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True, max_row=10))
    for r in rows:
        cell = str(r[1]) if len(r) > 1 and r[1] else ""
        if "Date Range:" in cell:
            match = re.search(r"(\d+/\d+/\d+) to (\d+/\d+/\d+)", cell)
            if match:
                start = datetime.strptime(match.group(1), "%m/%d/%Y")
                return start.strftime("%B %Y")
    return None


def compute_clinic_ltv(employees):
    """Same computation as the CLI's save_prev() clinic_ltv calc."""
    clinic_data = defaultdict(lambda: {'poor_guests': 0, 'poor_memb': 0})
    for e in employees:
        if e['close'] < POOR_THRESHOLD:
            clinic_data[e['clinic']]['poor_guests'] += e['guests']
            clinic_data[e['clinic']]['poor_memb']   += e['memb']
    return {c: max(0, round(d['poor_guests'] * 0.20) - d['poor_memb']) * LTV for c, d in clinic_data.items()}


def snapshot_for_history(employees, month_label, clinic_ltv):
    """Shaped like the old sales_dashboard_prev.json -- becomes `prev` on the next parse_report() call."""
    return {
        "month_label": month_label,
        "run_date": datetime.now().strftime("%Y-%m-%d"),
        "employees": {
            e["key"]: {"guests": e["guests"], "memb": e["memb"], "close": e["close"]}
            for e in employees
        },
        "clinic_ltv": clinic_ltv,
    }
